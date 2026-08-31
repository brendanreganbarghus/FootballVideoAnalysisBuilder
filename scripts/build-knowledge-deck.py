from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Football-Analytics-Knowledge-Session.pptx"

NAVY = RGBColor(13, 17, 23)
PANEL = RGBColor(22, 27, 34)
PANEL_2 = RGBColor(30, 38, 48)
WHITE = RGBColor(240, 246, 252)
MUTED = RGBColor(139, 148, 158)
BLUE = RGBColor(57, 160, 255)
GREEN = RGBColor(86, 211, 100)
ORANGE = RGBColor(255, 180, 84)
PURPLE = RGBColor(210, 168, 255)
RED = RGBColor(248, 81, 73)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: RGBColor = WHITE,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    font: str = "Aptos",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 17,
    color: RGBColor = WHITE,
    spacing: int = 7,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(spacing)
        properties = paragraph._p.get_or_add_pPr()
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "\u2022")
        properties.insert(0, bullet)
    return box


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor = PANEL,
    line: RGBColor = BLUE,
    radius: bool = True,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.4)
    return shape


def add_card(
    slide,
    title: str,
    detail: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    accent: RGBColor = BLUE,
    title_size: int = 17,
    detail_size: int = 13,
):
    add_rect(slide, x, y, w, h, line=accent)
    add_text(
        slide,
        title,
        x + 0.14,
        y + 0.12,
        w - 0.28,
        0.35,
        size=title_size,
        color=accent,
        bold=True,
    )
    add_text(
        slide,
        detail,
        x + 0.14,
        y + 0.52,
        w - 0.28,
        h - 0.62,
        size=detail_size,
        color=WHITE,
    )


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color=BLUE):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(2.25)
    line.line.end_arrowhead = True
    return line


def set_background(slide, color: RGBColor = NAVY):
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None):
    add_text(slide, title, 0.55, 0.3, 12.2, 0.55, size=27, bold=True)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.56),
        Inches(0.92),
        Inches(1.15),
        Inches(0.05),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 1.88, 0.77, 10.8, 0.28, size=12, color=MUTED)


def add_footer(slide, number: int):
    add_text(
        slide,
        "Football Video Analytics POC",
        0.55,
        7.15,
        4.2,
        0.2,
        size=9,
        color=MUTED,
    )
    add_text(
        slide,
        str(number),
        12.2,
        7.13,
        0.55,
        0.22,
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def new_slide(prs: Presentation, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def architecture_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Target home-ground architecture",
        "Two fixed views, local compute, no cloud dependency",
    )
    add_card(
        slide,
        "Camera A",
        "4K / 25 fps\nLeft-half fixed view\nWeatherproof tripod",
        0.55,
        1.25,
        2.1,
        1.35,
        accent=BLUE,
    )
    add_card(
        slide,
        "Camera B",
        "4K / 25 fps\nRight-half fixed view\nWeatherproof tripod",
        0.55,
        4.05,
        2.1,
        1.35,
        accent=GREEN,
    )
    add_card(
        slide,
        "Gigabit PoE switch",
        "Cat6 carries power + RTSP video\nLocal network; 100 m runs",
        3.25,
        2.55,
        2.2,
        1.4,
        accent=ORANGE,
        title_size=13,
        detail_size=11,
    )
    add_card(
        slide,
        "NVIDIA GPU workstation",
        "Decode both streams\nSynchronize timestamps\nDetect, track, fuse, infer",
        6.05,
        2.25,
        2.45,
        2.0,
        accent=PURPLE,
        title_size=13,
    )
    add_card(
        slide,
        "Operator review",
        "Confirm goals, shots on target\nCorrect disputed events",
        9.15,
        1.25,
        2.15,
        1.55,
        accent=ORANGE,
    )
    add_card(
        slide,
        "Stadium screen",
        "Fullscreen browser via HDMI\nProvisional AI Statistics",
        9.15,
        4.1,
        2.15,
        1.55,
        accent=GREEN,
    )
    add_arrow(slide, 2.65, 1.9, 3.25, 3.1, BLUE)
    add_arrow(slide, 2.65, 4.7, 3.25, 3.4, GREEN)
    add_arrow(slide, 5.45, 3.25, 6.05, 3.25, ORANGE)
    add_arrow(slide, 8.5, 2.95, 9.15, 1.95, PURPLE)
    add_arrow(slide, 8.5, 3.55, 9.15, 4.8, PURPLE)
    add_text(
        slide,
        "Why Cat6/PoE?",
        0.8,
        6.0,
        1.45,
        0.3,
        size=13,
        bold=True,
        color=BLUE,
    )
    add_text(
        slide,
        "Long outdoor runs, independent streams, camera power and standard network recovery. "
        "HDMI is best kept between the local PC and display.",
        2.25,
        6.0,
        10.15,
        0.42,
        size=12,
        color=MUTED,
    )


def process_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "End-to-end processing flow",
        "The same domain stages work offline today and on RTSP frames tomorrow",
    )
    stages = [
        ("1. Ingest", "RTSP decode\nTimestamp\nRecord originals", BLUE),
        ("2. Synchronize", "Pair A/B frames\nDetect drift\nDrop policy", ORANGE),
        ("3. Detect", "4K tiling\nYOLO player/ball\nOverlap NMS", PURPLE),
        ("4. Track", "Ball trajectory\nPlayer identities\nTeam colors", GREEN),
        ("5. Fuse", "Pitch calibration\nCross-camera IDs\nBest ball view", BLUE),
        ("6. Infer", "Possession\nPass/turnover\nShot direction", ORANGE),
        ("7. Publish", "Provisional totals\nOperator approval\nScreen UI", GREEN),
    ]
    x = 0.5
    for index, (title, detail, accent) in enumerate(stages):
        add_card(
            slide,
            title,
            detail,
            x,
            2.15,
            1.55,
            2.0,
            accent=accent,
            title_size=15,
            detail_size=12,
        )
        if index < len(stages) - 1:
            add_arrow(slide, x + 1.55, 3.15, x + 1.78, 3.15, MUTED)
        x += 1.8
    add_rect(slide, 1.0, 5.0, 11.25, 0.9, fill=PANEL_2, line=MUTED)
    add_text(
        slide,
        "Persistent state crosses frame and chunk boundaries: camera clocks, "
        "track IDs, ball history, possession owner, accepted events and "
        "finalized time.",
        1.25,
        5.27,
        10.75,
        0.4,
        size=16,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def package_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Software packages and responsibilities",
        "Python orchestrates; compiled C++ and CUDA perform the expensive work",
    )
    cards = [
        ("Python", "Pipeline, domain rules, CLIs, JSON and tests", BLUE),
        ("Ultralytics", "YOLO model loading, inference and class outputs", PURPLE),
        ("PyTorch / CUDA", "GPU tensor execution underneath YOLO", GREEN),
        ("OpenCV", "Decode, frame access, resize, draw and encode", ORANGE),
        ("NumPy", "Array and geometry operations", BLUE),
        ("LAP", "Linear assignment support for tracking", PURPLE),
        ("pytest", "Algorithm and regression verification", GREEN),
        ("Browser JS", "Video-synchronized counters and screen UI", ORANGE),
        ("python-pptx", "Rebuilds this editable knowledge deck", BLUE),
    ]
    for index, card in enumerate(cards):
        row, col = divmod(index, 3)
        add_card(
            slide,
            *card[:2],
            0.65 + col * 4.2,
            1.35 + row * 1.75,
            3.75,
            1.35,
            accent=card[2],
        )


def module_map_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Where the code lives",
        "Use function names as stable navigation anchors",
    )
    modules = [
        ("benchmark.py", "4K tiling, YOLO cache, NMS", BLUE),
        ("ball_tracking.py", "false-ball filtering, tracks, interpolation", GREEN),
        ("player_tracking.py", "identity association, team color, render", ORANGE),
        ("possession.py", "control, flight state, passes, shots, metrics", PURPLE),
        ("chunk_simulator.py", "overlap, deduplication, checkpoints, latency", BLUE),
        ("demo.py", "HTML dashboard and synchronized counters", GREEN),
        ("soccertrack.py", "ground-truth adapter and window selection", ORANGE),
        ("aggregate.py", "multi-window micro precision and recall", PURPLE),
    ]
    for index, (title, detail, accent) in enumerate(modules):
        row, col = divmod(index, 2)
        add_card(
            slide,
            title,
            detail,
            0.8 + col * 6.25,
            1.2 + row * 1.35,
            5.75,
            1.0,
            accent=accent,
            detail_size=12,
        )
    add_text(
        slide,
        "CLI wrappers: *_cli.py   |   Tests: tests\\test_<stage>.py   |   "
        "Full cached runner: scripts\\run-soccertrack-window.ps1",
        0.85,
        6.55,
        11.8,
        0.32,
        size=13,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def event_state_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Pass and turnover state machine",
        "possession.py: infer_flight_transfer_events + fallback deduplication",
    )
    states = [
        ("Stable owner", "Ball near player's feet\nTeam + track known", BLUE),
        ("Release", "Observed ball speed >= threshold\nNo synthetic velocity", ORANGE),
        ("Flight", "Follow accepted track\nPermit short missing gaps", PURPLE),
        ("First receiver", "Different player gains control\nWithin receiver window", GREEN),
        ("Classify", "Same team = pass\nOther team = turnover", BLUE),
    ]
    for index, (title, detail, accent) in enumerate(states):
        x = 0.55 + index * 2.55
        add_card(
            slide,
            title,
            detail,
            x,
            2.0,
            2.05,
            1.85,
            accent=accent,
            title_size=15,
            detail_size=12,
        )
        if index < len(states) - 1:
            add_arrow(slide, x + 2.05, 2.92, x + 2.47, 2.92, MUTED)
    add_card(
        slide,
        "Stable-control fallback",
        "Fallback recovers events when flight evidence is incomplete.\n"
        "Nearby duplicates are suppressed before output.",
        1.55,
        4.65,
        4.9,
        1.15,
        accent=ORANGE,
        detail_size=11,
    )
    add_card(
        slide,
        "Current failure modes",
        "Tiny ball, motion blur, one-touch sequences, identity switches,\n"
        "team-color noise and missing receiver control.",
        6.9,
        4.65,
        4.9,
        1.15,
        accent=RED,
        detail_size=11,
    )


def shot_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Shot inference and future outcomes",
        "possession.py: infer_shot_events",
    )
    add_card(
        slide,
        "Evidence used today",
        "Observed ball speed\nDirection toward calibrated goal center\nNo quick "
        "receiver control\nLatest confirmed possession owner",
        0.8,
        1.35,
        3.55,
        2.4,
        accent=ORANGE,
    )
    add_card(
        slide,
        "Can classify today",
        "Shot candidate\nLikely attacking team\nTimestamp and confidence",
        4.9,
        1.35,
        3.55,
        2.4,
        accent=GREEN,
    )
    add_card(
        slide,
        "Needs goal-camera evidence",
        "Goal-plane crossing\nShot on target\nGoalkeeper save\nMiss / post / "
        "crossbar",
        9.0,
        1.35,
        3.55,
        2.4,
        accent=RED,
    )
    add_arrow(slide, 4.35, 2.55, 4.9, 2.55, MUTED)
    add_arrow(slide, 8.45, 2.55, 9.0, 2.55, MUTED)
    add_rect(slide, 1.25, 4.7, 10.8, 1.0, fill=PANEL_2, line=PURPLE)
    add_text(
        slide,
        "Two behind-goal cameras are a later phase. The first home pilot uses "
        "two sideline half-pitch views and operator confirmation for public "
        "shot outcomes.",
        1.5,
        4.98,
        10.3,
        0.45,
        size=16,
        align=PP_ALIGN.CENTER,
        bold=True,
    )


def tuning_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Main tuning surfaces",
        "Change one behavior at a time and evaluate all benchmark windows",
    )
    add_card(
        slide,
        "Detection",
        "--confidence\n--image-size\n--stride\n--tile-width\n--overlap\n--nms-iou",
        0.55,
        1.3,
        2.75,
        3.7,
        accent=BLUE,
    )
    add_card(
        slide,
        "Ball tracking",
        "static occupancy\nmax gap\nmax speed\nminimum points\ninterpolation "
        "constraints",
        3.55,
        1.3,
        2.75,
        3.7,
        accent=GREEN,
    )
    add_card(
        slide,
        "Player/team",
        "association gap\nplayer speed\nminimum points\njersey feature ranges\n"
        "team smoothing",
        6.55,
        1.3,
        2.75,
        3.7,
        accent=ORANGE,
    )
    add_card(
        slide,
        "Events",
        "control radius\npass speed\nreceiver window\ndebounce / dedup\nshot "
        "speed + cosine",
        9.55,
        1.3,
        2.75,
        3.7,
        accent=PURPLE,
    )
    add_text(
        slide,
        "Rule: observed points establish kick/shot velocity. Interpolated points "
        "may support possession continuity but never invent an event trajectory.",
        1.1,
        5.65,
        11.1,
        0.55,
        size=15,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def commands_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Developer edit-test-run loop",
        "Editable install means source changes are active on the next command",
    )
    commands = (
        'Setup\n'
        'python -m venv .venv\n'
        '.\\.venv\\Scripts\\python -m pip install -e ".[test,docs,dataset]"\n\n'
        'Validate\n'
        'python -m compileall -q src tests\n'
        'python -m pytest -q\n'
        'python -m pip check\n\n'
        'Refresh event logic and demo\n'
        'python -m football_poc.possession_cli\n'
        'python -m football_poc.chunk_simulator_cli\n'
        'python -m football_poc.demo_cli\n\n'
        'Serve\n'
        'python -m http.server 8000'
    )
    add_rect(slide, 0.7, 1.25, 7.25, 5.45, fill=RGBColor(9, 12, 16), line=BLUE)
    add_text(
        slide,
        commands,
        0.95,
        1.5,
        6.75,
        4.95,
        size=12,
        color=RGBColor(165, 214, 255),
        font="Cascadia Mono",
    )
    add_card(
        slide,
        "Rerun only downstream stages",
        "Event change: possession -> chunk -> aggregate -> demo\n\n"
        "Tracking change: tracking -> possession -> downstream\n\n"
        "Detector change: new cache -> every downstream stage",
        8.35,
        1.25,
        4.25,
        3.0,
        accent=ORANGE,
        detail_size=11,
    )
    add_card(
        slide,
        "Full reference",
        "docs\\DEVELOPER_GUIDE.md\n\n"
        "scripts\\run-soccertrack-window.ps1\n\n"
        "scripts\\build-knowledge-deck.py",
        8.35,
        4.6,
        4.25,
        1.75,
        accent=GREEN,
    )


def results_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "What has been achieved",
        "Measured POC evidence, not a production accuracy claim",
    )
    metrics = [
        ("45.9%", "Tracked-frame ball coverage\nafter conservative interpolation", BLUE),
        ("75.0%", "Pass precision\nfirst benchmark minute", GREEN),
        ("63.2%", "Pass recall\nfirst benchmark minute", ORANGE),
        ("71.0%", "Pass precision\nacross three windows", GREEN),
        ("47.8%", "Pass recall\nacross three windows", ORANGE),
        ("0.4 fps", "Current CPU panoramic throughput\n~31x below target", RED),
    ]
    for index, (value, detail, accent) in enumerate(metrics):
        row, col = divmod(index, 3)
        x = 0.65 + col * 4.2
        y = 1.35 + row * 2.35
        add_rect(slide, x, y, 3.75, 1.9, line=accent)
        add_text(
            slide,
            value,
            x + 0.15,
            y + 0.18,
            3.45,
            0.62,
            size=30,
            color=accent,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            detail,
            x + 0.2,
            y + 0.92,
            3.35,
            0.7,
            size=14,
            align=PP_ALIGN.CENTER,
        )


def roadmap_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Roadmap to the home-match pilot",
        "Evidence-driven gates before public stadium use",
    )
    steps = [
        ("1", "Camera proof", "Record one original minute; measure ball pixels, blur, bitrate."),
        ("2", "GPU benchmark", "Measure two-stream decode + inference throughput and latency."),
        ("3", "RTSP ingest", "Reconnect, timestamp, synchronize, record and monitor both feeds."),
        ("4", "Camera fusion", "Calibrate both views into common pitch coordinates."),
        ("5", "Operator UI", "Approve/correct sensitive events before public display."),
        ("6", "Closed pilot", "Run complete home match without showing public statistics."),
        ("7", "Screen pilot", "Show limited provisional metrics at 10-20 s latency."),
    ]
    for index, (number, title, detail) in enumerate(steps):
        y = 1.15 + index * 0.76
        add_rect(slide, 0.8, y, 0.58, 0.58, fill=BLUE, line=BLUE)
        add_text(
            slide,
            number,
            0.8,
            y + 0.03,
            0.58,
            0.45,
            size=16,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, title, 1.65, y + 0.03, 2.2, 0.35, size=17, bold=True)
        add_text(slide, detail, 3.85, y + 0.04, 8.4, 0.4, size=14, color=MUTED)


def knowledge_session_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Suggested knowledge-session flow",
        "A 35-45 minute technical walkthrough",
    )
    agenda = [
        ("5 min", "Problem and amateur-club opportunity", BLUE),
        ("5 min", "Current demo and honest benchmark results", GREEN),
        ("8 min", "Detection, tracking and event state machine", PURPLE),
        ("7 min", "Two-camera local-GPU target architecture", ORANGE),
        ("8 min", "Repository tour and live CLI demonstration", BLUE),
        ("5 min", "Limitations, licensing, privacy and roadmap", RED),
        ("5 min", "Questions and camera/GPU pilot decision", GREEN),
    ]
    for index, (time, topic, accent) in enumerate(agenda):
        y = 1.2 + index * 0.78
        add_rect(slide, 0.95, y, 1.05, 0.5, fill=PANEL_2, line=accent)
        add_text(
            slide,
            time,
            0.98,
            y + 0.07,
            0.98,
            0.3,
            size=14,
            color=accent,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, topic, 2.3, y + 0.04, 9.5, 0.4, size=17)
    add_text(
        slide,
        "Recommended live commands: focused pytest -> possession refresh -> "
        "demo refresh -> browser playback.",
        2.3,
        6.72,
        9.5,
        0.28,
        size=12,
        color=MUTED,
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Football Analytics Knowledge Session"
    prs.core_properties.subject = (
        "Current POC, code structure and two-camera local GPU architecture"
    )
    prs.core_properties.author = "Football Video Analytics Builder"

    title = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(title)
    add_text(
        title,
        "AI Football Analytics",
        0.8,
        1.35,
        11.7,
        0.85,
        size=38,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        title,
        "Knowledge session: from 4K video to live amateur-club statistics",
        1.15,
        2.35,
        11.0,
        0.55,
        size=22,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_rect(title, 2.25, 3.35, 8.85, 1.35, fill=PANEL, line=GREEN)
    add_text(
        title,
        "Current POC + code map + two-camera Cat6/PoE local-GPU roadmap",
        2.55,
        3.72,
        8.25,
        0.6,
        size=20,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        title,
        "FootballVideoAnalysisBuilder",
        4.25,
        6.25,
        4.85,
        0.3,
        size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    slide = new_slide(prs, "The problem we are solving")
    add_bullets(
        slide,
        [
            "Amateur clubs rarely have affordable, near-live match statistics.",
            "Manual analysts are expensive and cannot scale across every match.",
            "A distant full-pitch view makes the ball only a few pixels wide.",
            "The product goal is provisional, reviewable statistics on the home-ground screen.",
        ],
        0.8,
        1.35,
        6.0,
        4.6,
        size=20,
        spacing=15,
    )
    add_card(
        slide,
        "Product distinction",
        "Affordable and transparent local processing using the club's own fixed "
        "cameras - not professional league tracking infrastructure.",
        7.35,
        1.55,
        4.9,
        1.65,
        accent=GREEN,
        title_size=20,
        detail_size=16,
    )
    add_card(
        slide,
        "Public-screen principle",
        "Candidate events remain provisional until confidence or an operator "
        "review allows publication.",
        7.35,
        3.65,
        4.9,
        1.65,
        accent=ORANGE,
        title_size=20,
        detail_size=16,
    )

    results_slide(prs)

    slide = new_slide(
        prs,
        "Current POC versus target system",
        "Be explicit about what works today and what is next",
    )
    add_card(
        slide,
        "Implemented today",
        "Recorded single panorama\nResumable tiled detection cache\nBall and "
        "player tracking\nTeam colors\nPass/turnover/shot candidates\nThree-window "
        "evaluation\nChunk replay and demo",
        0.8,
        1.3,
        5.55,
        4.8,
        accent=GREEN,
        title_size=21,
        detail_size=17,
    )
    add_card(
        slide,
        "Target home-ground system",
        "Two live RTSP cameras\nTimestamp synchronization\nLocal NVIDIA GPU\n"
        "Cross-camera pitch fusion\nPersistent live state\nOperator review\n"
        "WebSocket statistics\nStadium screen output",
        6.95,
        1.3,
        5.55,
        4.8,
        accent=BLUE,
        title_size=21,
        detail_size=17,
    )

    architecture_slide(prs)
    process_slide(prs)
    package_slide(prs)
    module_map_slide(prs)

    slide = new_slide(
        prs,
        "Detection and cache stage",
        "benchmark.py - preserve expensive inference so algorithms can iterate quickly",
    )
    add_card(
        slide,
        "4K panorama",
        "4096 x 1080 source\n25 fps\nExact manifest frame window",
        0.7,
        1.55,
        2.2,
        1.55,
        accent=BLUE,
    )
    add_arrow(slide, 2.9, 2.32, 3.3, 2.32, MUTED)
    add_card(
        slide,
        "Horizontal tiles",
        "1280 px tiles\n20% overlap\nRetain small-ball detail",
        3.3,
        1.55,
        2.2,
        1.55,
        accent=ORANGE,
    )
    add_arrow(slide, 5.5, 2.32, 5.9, 2.32, MUTED)
    add_card(
        slide,
        "YOLO inference",
        "Football classes\nconfidence 0.01\nimage size 960",
        5.9,
        1.55,
        2.2,
        1.55,
        accent=PURPLE,
    )
    add_arrow(slide, 8.1, 2.32, 8.5, 2.32, MUTED)
    add_card(
        slide,
        "Class-aware NMS",
        "Offset tile boxes\nRemove overlap duplicates",
        8.5,
        1.55,
        2.2,
        1.55,
        accent=GREEN,
    )
    add_arrow(slide, 10.7, 2.32, 11.1, 2.32, MUTED)
    add_card(
        slide,
        "JSONL cache",
        "Append per frame\nResume safely\nReuse downstream",
        11.1,
        1.55,
        1.65,
        1.55,
        accent=BLUE,
        title_size=14,
        detail_size=11,
    )
    add_card(
        slide,
        "Why cache?",
        "Current CPU throughput is ~0.4 panoramas/s. Event and tracking changes "
        "must not trigger another 23-hour half-match inference run.",
        1.35,
        4.35,
        4.8,
        1.35,
        accent=RED,
    )
    add_card(
        slide,
        "GPU migration",
        "Start with PyTorch CUDA. Then benchmark ONNX/TensorRT or DeepStream only "
        "if measured throughput requires it.",
        7.0,
        4.35,
        4.8,
        1.35,
        accent=GREEN,
    )

    slide = new_slide(
        prs,
        "Tracking and team classification",
        "ball_tracking.py + player_tracking.py",
    )
    add_card(
        slide,
        "Ball pipeline",
        "Pitch mask\nStatic false-positive cells\nTime/speed association\nShort "
        "gap interpolation\nObserved/interpolated provenance",
        0.8,
        1.35,
        5.45,
        3.9,
        accent=GREEN,
        title_size=21,
        detail_size=17,
    )
    add_card(
        slide,
        "Player pipeline",
        "Pitch filter\nIdentity association\nTorso color features\nBlue/white/"
        "keeper/official\nPer-frame team labels\nVerification render",
        7.05,
        1.35,
        5.45,
        3.9,
        accent=ORANGE,
        title_size=21,
        detail_size=17,
    )
    add_text(
        slide,
        "Both feed possession.py; neither independently proves a pass or shot.",
        1.55,
        5.85,
        10.25,
        0.45,
        size=18,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    event_state_slide(prs)
    shot_slide(prs)

    slide = new_slide(
        prs,
        "Live-style chunk state",
        "chunk_simulator.py validates delivery behavior, not live model inference",
    )
    add_card(
        slide,
        "20 s chunk",
        "Arrives with 2 s overlap\nContains precomputed candidate events",
        0.85,
        1.45,
        2.45,
        1.6,
        accent=BLUE,
    )
    add_arrow(slide, 3.3, 2.25, 3.75, 2.25, MUTED)
    add_card(
        slide,
        "Deduplicate",
        "Event type + timestamp + team + track IDs",
        3.75,
        1.45,
        2.45,
        1.6,
        accent=ORANGE,
    )
    add_arrow(slide, 6.2, 2.25, 6.65, 2.25, MUTED)
    add_card(
        slide,
        "Checkpoint",
        "Accepted events\nNext chunk\nWorker availability",
        6.65,
        1.45,
        2.45,
        1.6,
        accent=PURPLE,
    )
    add_arrow(slide, 9.1, 2.25, 9.55, 2.25, MUTED)
    add_card(
        slide,
        "Publish",
        "Provisional totals\nFinalized-through time\nQueue latency",
        9.55,
        1.45,
        2.45,
        1.6,
        accent=GREEN,
    )
    add_card(
        slide,
        "Future live replacement",
        "The RTSP service must carry the same state continuously across frames "
        "and reconnects. It should not restart identities at every chunk.",
        1.2,
        4.35,
        4.95,
        1.35,
        accent=BLUE,
    )
    add_card(
        slide,
        "Latency condition",
        "With 20 s chunks and 2 s overlap, processing must complete faster than "
        "the 18 s arrival cadence to avoid growing backlog.",
        7.15,
        4.35,
        4.95,
        1.35,
        accent=RED,
    )

    tuning_slide(prs)
    commands_slide(prs)

    slide = new_slide(
        prs,
        "Outputs and evidence",
        "Every stage leaves an inspectable artifact",
    )
    outputs = [
        ("detections.jsonl", "Raw cached player/ball boxes per sampled frame", BLUE),
        ("ball-tracks.json", "Accepted trajectories with interpolation flags", GREEN),
        ("player-tracks.json", "Identities, boxes and per-frame team labels", ORANGE),
        ("possession.json", "Control observations and stable segments", PURPLE),
        ("predicted-events.json", "Pass, turnover and shot candidates", BLUE),
        ("event-evaluation.json", "Timestamp matches, precision and recall", GREEN),
        ("chunk-simulation.json", "Provisional state, dedup and latency", ORANGE),
        ("demo\\index.html", "Team-facing synchronized replay", PURPLE),
    ]
    for index, (title, detail, accent) in enumerate(outputs):
        row, col = divmod(index, 2)
        add_card(
            slide,
            title,
            detail,
            0.75 + col * 6.3,
            1.2 + row * 1.32,
            5.8,
            0.98,
            accent=accent,
            detail_size=12,
        )

    slide = new_slide(
        prs,
        "Engineering and product guardrails",
        "Accuracy, licensing and privacy remain first-class work",
    )
    add_bullets(
        slide,
        [
            "The football checkpoint has unresolved license/training-data provenance; internal POC only.",
            "Ultralytics distribution has AGPL/commercial licensing implications.",
            "Public statistics must distinguish candidates from confirmed events.",
            "Minors, consent, retention, deletion and processor agreements require a DPIA/AVG plan.",
            "Camera failures, GPU backlog and dropped frames must be visible to the operator.",
            "Thresholds must generalize across held-out matches, weather and floodlights.",
        ],
        0.85,
        1.3,
        11.55,
        4.8,
        size=18,
        spacing=11,
    )

    roadmap_slide(prs)
    knowledge_session_slide(prs)

    slide = new_slide(prs, "Key takeaways")
    add_card(
        slide,
        "The POC is real",
        "It produces inspectable tracks, team labels, event candidates, benchmark "
        "metrics and synchronized live-style counters.",
        0.9,
        1.35,
        3.55,
        2.1,
        accent=GREEN,
        title_size=21,
        detail_size=16,
    )
    add_card(
        slide,
        "Known bottleneck",
        "Tiny-ball evidence and CPU throughput - not the choice of Python as an "
        "orchestration language.",
        4.9,
        1.35,
        3.55,
        2.1,
        accent=ORANGE,
        title_size=21,
        detail_size=16,
    )
    add_card(
        slide,
        "Measurable next proof",
        "One original camera minute, two-stream GPU benchmark, then a closed "
        "home-match pilot before public display.",
        8.9,
        1.35,
        3.55,
        2.1,
        accent=BLUE,
        title_size=21,
        detail_size=16,
    )
    add_text(
        slide,
        "Detailed code navigation and commands: docs\\DEVELOPER_GUIDE.md",
        1.1,
        5.2,
        11.1,
        0.55,
        size=20,
        color=PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"Knowledge deck written to {build().resolve()}")
