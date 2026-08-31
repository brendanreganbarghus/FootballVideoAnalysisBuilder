from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Football-Analytics-Product-Pitch.pptx"

NAVY = RGBColor(10, 18, 28)
PANEL = RGBColor(20, 31, 43)
PANEL_2 = RGBColor(29, 44, 58)
WHITE = RGBColor(245, 249, 252)
MUTED = RGBColor(157, 170, 182)
BLUE = RGBColor(50, 166, 255)
GREEN = RGBColor(72, 210, 112)
PITCH = RGBColor(31, 130, 73)
PITCH_DARK = RGBColor(23, 103, 57)
ORANGE = RGBColor(255, 179, 71)
PURPLE = RGBColor(195, 148, 255)
RED = RGBColor(255, 91, 91)
BLACK = RGBColor(4, 8, 12)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: RGBColor = WHITE,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_shape(
    slide,
    shape_type,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor,
    line: RGBColor | None = None,
    width: float = 1.2,
):
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: RGBColor = WHITE,
    width: float = 1.5,
    arrow: bool = False,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def set_background(slide, color: RGBColor = NAVY):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_footer(slide, number: int):
    add_text(
        slide,
        "Affordable local football analytics",
        0.55,
        7.12,
        4.5,
        0.2,
        size=9,
        color=MUTED,
    )
    add_text(
        slide,
        str(number),
        12.2,
        7.12,
        0.55,
        0.2,
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_title(slide, title: str, subtitle: str | None = None):
    add_text(slide, title, 0.58, 0.28, 12.1, 0.55, size=28, bold=True)
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        0.58,
        0.9,
        1.05,
        0.05,
        fill=GREEN,
    )
    if subtitle:
        add_text(slide, subtitle, 1.85, 0.76, 10.8, 0.25, size=12, color=MUTED)


def new_slide(prs: Presentation, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def add_card(
    slide,
    title: str,
    detail: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    accent: RGBColor = GREEN,
    title_size: int = 20,
    detail_size: int = 15,
):
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=PANEL,
        line=accent,
        width=1.4,
    )
    add_text(
        slide,
        title,
        x + 0.18,
        y + 0.16,
        w - 0.36,
        0.38,
        size=title_size,
        color=accent,
        bold=True,
    )
    add_text(
        slide,
        detail,
        x + 0.18,
        y + 0.65,
        w - 0.36,
        h - 0.78,
        size=detail_size,
    )


def add_badge(slide, text: str, x: float, y: float, w: float, color: RGBColor):
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        0.38,
        fill=PANEL_2,
        line=color,
    )
    add_text(
        slide,
        text,
        x + 0.05,
        y + 0.06,
        w - 0.1,
        0.22,
        size=11,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_pitch(slide, x: float, y: float, w: float, h: float):
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=PITCH,
        line=WHITE,
        width=1.6,
    )
    for stripe in range(1, 6, 2):
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            x + stripe * w / 6,
            y + 0.02,
            w / 6,
            h - 0.04,
            fill=PITCH_DARK,
        )
    add_line(slide, x + w / 2, y, x + w / 2, y + h, width=1.1)
    centre = add_shape(
        slide,
        MSO_SHAPE.OVAL,
        x + w / 2 - 0.48,
        y + h / 2 - 0.48,
        0.96,
        0.96,
        fill=PITCH,
        line=WHITE,
    )
    centre.fill.transparency = 100000
    add_shape(
        slide,
        MSO_SHAPE.OVAL,
        x + w / 2 - 0.05,
        y + h / 2 - 0.05,
        0.1,
        0.1,
        fill=WHITE,
    )
    for side_x in (x + 0.02, x + w - 1.37):
        area = add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            side_x,
            y + h * 0.23,
            1.35,
            h * 0.54,
            fill=PITCH,
            line=WHITE,
        )
        area.fill.transparency = 100000
    for player_x, player_y, color in (
        (x + 2.0, y + 1.0, BLUE),
        (x + 3.0, y + 2.0, BLUE),
        (x + 4.2, y + 0.8, BLUE),
        (x + 4.7, y + 2.5, BLUE),
        (x + w - 2.0, y + 1.0, ORANGE),
        (x + w - 3.0, y + 2.1, ORANGE),
        (x + w - 4.1, y + 0.75, ORANGE),
        (x + w - 4.7, y + 2.55, ORANGE),
    ):
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            player_x,
            player_y,
            0.18,
            0.18,
            fill=color,
            line=WHITE,
            width=0.7,
        )
    add_shape(
        slide,
        MSO_SHAPE.OVAL,
        x + w / 2 + 0.35,
        y + h / 2 - 0.1,
        0.12,
        0.12,
        fill=WHITE,
        line=BLACK,
        width=0.7,
    )


def add_camera_tower(
    slide,
    x: float,
    y: float,
    *,
    facing: str,
    label: str,
    accent: RGBColor,
):
    add_line(slide, x + 0.45, y + 0.7, x + 0.45, y + 2.22, color=MUTED, width=3)
    add_line(slide, x + 0.45, y + 1.78, x + 0.1, y + 2.28, color=MUTED, width=2)
    add_line(slide, x + 0.45, y + 1.78, x + 0.8, y + 2.28, color=MUTED, width=2)
    camera_x = x + (0.13 if facing == "right" else 0.27)
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        camera_x,
        y + 0.4,
        0.55,
        0.36,
        fill=PANEL_2,
        line=accent,
    )
    lens_x = camera_x + (0.46 if facing == "right" else -0.06)
    add_shape(
        slide,
        MSO_SHAPE.OVAL,
        lens_x,
        y + 0.49,
        0.16,
        0.16,
        fill=accent,
        line=WHITE,
        width=0.6,
    )
    add_text(
        slide,
        label,
        x - 0.2,
        y,
        1.3,
        0.35,
        size=13,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "6–10 m",
        x - 0.22,
        y + 1.15,
        0.6,
        0.3,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(
        slide,
        x - 0.05,
        y + 0.72,
        x - 0.05,
        y + 2.18,
        color=WHITE,
        width=1,
    )


def add_sideline_camera(
    slide,
    x: float,
    y: float,
    *,
    direction: str,
    label: str,
    accent: RGBColor,
):
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        0.7,
        0.34,
        fill=PANEL_2,
        line=accent,
        width=1.4,
    )
    lens_y = y + 0.25 if direction == "down" else y - 0.07
    add_shape(
        slide,
        MSO_SHAPE.OVAL,
        x + 0.27,
        lens_y,
        0.17,
        0.17,
        fill=accent,
        line=WHITE,
        width=0.6,
    )
    label_y = y - 0.28 if direction == "down" else y + 0.39
    add_text(
        slide,
        label,
        x - 0.55,
        label_y,
        1.8,
        0.24,
        size=11,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_gpu(slide, x: float, y: float):
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        1.7,
        1.15,
        fill=PANEL,
        line=PURPLE,
        width=1.6,
    )
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        x + 0.16,
        y + 0.18,
        0.55,
        0.75,
        fill=BLACK,
        line=MUTED,
    )
    for fan_y in (y + 0.29, y + 0.6):
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            x + 0.31,
            fan_y,
            0.22,
            0.22,
            fill=PANEL_2,
            line=PURPLE,
        )
    add_text(
        slide,
        "LOCAL\nGPU PC",
        x + 0.8,
        y + 0.26,
        0.76,
        0.55,
        size=14,
        color=PURPLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_scoreboard(slide, x: float, y: float, w: float = 2.2, h: float = 1.25):
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        x,
        y,
        w,
        h,
        fill=BLACK,
        line=GREEN,
        width=2,
    )
    add_text(
        slide,
        "PROVISIONAL AI STATS",
        x + 0.08,
        y + 0.08,
        w - 0.16,
        0.24,
        size=11,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "BLUE  12 PASSES  1 SHOT\nWHITE   8 PASSES  0 SHOTS",
        x + 0.12,
        y + 0.43,
        w - 0.24,
        0.55,
        size=13,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, x + 0.25, y + h, x + 0.15, y + h + 0.32, color=MUTED, width=2)
    add_line(
        slide,
        x + w - 0.25,
        y + h,
        x + w - 0.15,
        y + h + 0.32,
        color=MUTED,
        width=2,
    )


def cover_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_badge(slide, "PRODUCT VISION", 0.72, 0.55, 1.55, GREEN)
    add_text(
        slide,
        "Live football statistics\nfor every club",
        0.72,
        1.15,
        6.1,
        1.65,
        size=38,
        bold=True,
    )
    add_text(
        slide,
        "Affordable cameras. Local AI processing. Provisional match statistics "
        "on the stadium screen.",
        0.75,
        3.03,
        5.65,
        1.0,
        size=21,
        color=MUTED,
    )
    add_badge(slide, "NO MANDATORY CLOUD", 0.75, 4.35, 1.9, BLUE)
    add_badge(slide, "AMATEUR-CLUB FOCUS", 2.83, 4.35, 2.0, ORANGE)
    add_badge(slide, "HUMAN REVIEW", 5.0, 4.35, 1.55, PURPLE)
    add_pitch(slide, 7.05, 1.25, 5.25, 3.35)
    add_sideline_camera(
        slide,
        8.25,
        1.02,
        direction="down",
        label="CAM A • TOUCHLINE",
        accent=BLUE,
    )
    add_sideline_camera(
        slide,
        10.5,
        4.48,
        direction="up",
        label="CAM B • TOUCHLINE",
        accent=ORANGE,
    )
    add_text(
        slide,
        "FootballVideoAnalysisBuilder",
        0.75,
        6.85,
        4.2,
        0.25,
        size=10,
        color=MUTED,
    )


def market_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "The market already proves the need",
        "Professional-quality systems exist—but usually as closed camera and cloud platforms",
    )
    add_card(
        slide,
        "What exists today",
        "Veo, Pixellot, Spiideo and Zone14 combine wide-angle cameras, AI tracking "
        "and online analysis.",
        0.7,
        1.35,
        3.75,
        3.1,
        accent=BLUE,
    )
    add_card(
        slide,
        "Why clubs buy them",
        "Less manual filming\nFaster review\nHighlights and tactical tools\nSome "
        "tracking and performance data",
        4.8,
        1.35,
        3.75,
        3.1,
        accent=GREEN,
    )
    add_card(
        slide,
        "The amateur gap",
        "Hardware lock-in\nCloud subscriptions\nEnterprise integrations\nCosts "
        "that are difficult for small clubs to sustain",
        8.9,
        1.35,
        3.75,
        3.1,
        accent=ORANGE,
    )
    add_text(
        slide,
        "The opportunity is not inventing video analytics—it is making useful "
        "analytics affordable, local and controllable.",
        1.25,
        5.15,
        10.8,
        0.8,
        size=24,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Market references: Veo, Pixellot, Spiideo and Zone14 public product pages.",
        2.2,
        6.25,
        8.9,
        0.25,
        size=10,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def opportunity_slide(prs: Presentation):
    slide = new_slide(prs, "Who we are building for")
    add_card(
        slide,
        "The coach",
        "Wants passes, possession changes, shots and match moments without assigning "
        "a person to count everything.",
        0.75,
        1.35,
        3.7,
        2.55,
        accent=BLUE,
    )
    add_card(
        slide,
        "The club",
        "Wants to own its cameras, footage and processing costs instead of depending "
        "on a costly external platform.",
        4.82,
        1.35,
        3.7,
        2.55,
        accent=GREEN,
    )
    add_card(
        slide,
        "The supporter",
        "Sees provisional live statistics on the home-ground screen and gets a "
        "more engaging match-day experience.",
        8.89,
        1.35,
        3.7,
        2.55,
        accent=ORANGE,
    )
    add_text(
        slide,
        "One locally owned system can support coaching, operations and the crowd.",
        1.3,
        4.75,
        10.7,
        0.65,
        size=27,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Initial focus: home matches, limited provisional statistics and operator review.",
        2.0,
        5.65,
        9.3,
        0.45,
        size=17,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def proposition_slide(prs: Presentation):
    slide = new_slide(prs, "Our product idea")
    add_text(
        slide,
        "Two fixed cameras watch the match.",
        0.9,
        1.4,
        5.1,
        0.55,
        size=28,
        bold=True,
    )
    add_text(
        slide,
        "A local GPU computer follows players and the ball.",
        0.9,
        2.25,
        5.1,
        0.75,
        size=28,
        bold=True,
        color=PURPLE,
    )
    add_text(
        slide,
        "The club receives reviewable statistics within seconds.",
        0.9,
        3.3,
        5.25,
        0.75,
        size=28,
        bold=True,
        color=GREEN,
    )
    add_scoreboard(slide, 7.45, 1.55, 4.25, 2.25)
    add_badge(slide, "LOCAL", 7.62, 4.3, 1.15, GREEN)
    add_badge(slide, "REVIEWABLE", 8.95, 4.3, 1.55, PURPLE)
    add_badge(slide, "LOWER RUNNING COST", 10.68, 4.3, 1.85, ORANGE)
    add_text(
        slide,
        "Product promise: useful rather than perfect, transparent rather than "
        "black-box, and designed around an amateur club’s budget.",
        1.2,
        5.55,
        10.9,
        0.8,
        size=21,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def setup_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "A simple home-match setup",
        "Two elevated fixed views feed one local computer—no cloud is required during the match",
    )
    add_pitch(slide, 0.72, 1.5, 8.7, 4.65)
    add_sideline_camera(
        slide,
        3.15,
        1.27,
        direction="down",
        label="CAM A • TOUCHLINE",
        accent=BLUE,
    )
    add_sideline_camera(
        slide,
        6.55,
        5.58,
        direction="up",
        label="CAM B • TOUCHLINE",
        accent=ORANGE,
    )
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        10.0,
        1.55,
        2.1,
        0.7,
        fill=PANEL,
        line=ORANGE,
    )
    add_text(
        slide,
        "PoE SWITCH",
        10.18,
        1.76,
        1.72,
        0.25,
        size=13,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_gpu(slide, 10.2, 2.65)
    add_scoreboard(slide, 9.78, 4.35, 2.55, 1.35)
    add_line(slide, 3.5, 1.37, 10.0, 1.9, color=BLUE, width=2.6, arrow=True)
    add_line(slide, 6.9, 5.68, 10.0, 1.9, color=ORANGE, width=2.6, arrow=True)
    add_line(slide, 11.05, 2.25, 11.05, 2.65, color=PURPLE, width=2.6, arrow=True)
    add_line(slide, 11.05, 3.8, 11.05, 4.35, color=GREEN, width=2.6, arrow=True)
    add_text(
        slide,
        "Cat6 / PoE",
        7.2,
        1.45,
        1.1,
        0.25,
        size=11,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Cat6 / PoE",
        8.2,
        4.55,
        1.1,
        0.25,
        size=11,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "HDMI / browser",
        11.2,
        4.0,
        1.35,
        0.25,
        size=11,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Both cameras mount 6–10 m above ground on stable, weatherproof "
        "tripods or masts beside the long touchlines.",
        0.95,
        6.45,
        8.2,
        0.42,
        size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_badge(slide, "BEHIND-GOAL CAMERAS: LATER", 9.62, 6.3, 2.95, PURPLE)


def process_slide(prs: Presentation):
    slide = new_slide(prs, "From the match to the screen")
    steps = [
        ("1", "Capture", "Both cameras record the whole playing area.", BLUE),
        ("2", "Understand", "The local PC follows players, teams and the ball.", PURPLE),
        ("3", "Calculate", "Football rules turn movement into candidate events.", ORANGE),
        ("4", "Publish", "An operator reviews important events before display.", GREEN),
    ]
    for index, (number, title, detail, color) in enumerate(steps):
        x = 0.65 + index * 3.15
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            x + 0.9,
            1.28,
            0.72,
            0.72,
            fill=color,
        )
        add_text(
            slide,
            number,
            x + 0.9,
            1.43,
            0.72,
            0.3,
            size=20,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_card(
            slide,
            title,
            detail,
            x,
            2.25,
            2.55,
            2.45,
            accent=color,
            title_size=22,
            detail_size=16,
        )
        if index < 3:
            add_line(
                slide,
                x + 2.58,
                3.47,
                x + 3.05,
                3.47,
                color=MUTED,
                width=2.2,
                arrow=True,
            )
    add_text(
        slide,
        "Target public latency: approximately 10–20 seconds for provisional statistics.",
        1.3,
        5.45,
        10.7,
        0.55,
        size=20,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def output_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "What the club and crowd will see",
        "Keep the public view simple; keep deeper evidence available for the operator",
    )
    add_scoreboard(slide, 0.85, 1.45, 5.15, 2.75)
    add_card(
        slide,
        "Live stadium view",
        "Passes\nShots\nTurnovers\nPossession trend\nLatest confirmed event",
        6.55,
        1.45,
        2.75,
        3.05,
        accent=GREEN,
        title_size=20,
        detail_size=16,
    )
    add_card(
        slide,
        "Operator view",
        "Video evidence\nAccept or correct\nAdd missed events\nMonitor camera and GPU health",
        9.65,
        1.45,
        2.75,
        3.05,
        accent=PURPLE,
        title_size=20,
        detail_size=16,
    )
    add_text(
        slide,
        "The public screen says “Provisional AI Statistics”—not “official match data.”",
        1.05,
        5.25,
        11.2,
        0.65,
        size=23,
        bold=True,
        color=ORANGE,
        align=PP_ALIGN.CENTER,
    )


def own_code_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "Why we are writing our own software",
        "Open components help—but the football decision-making and reliable product still have to be built",
    )
    items = [
        (
            "Control cost",
            "No compulsory per-match cloud processing or vendor camera lock-in.",
            GREEN,
        ),
        (
            "Tune for our cameras",
            "Train and calibrate for the exact height, lens, pitch and lighting.",
            BLUE,
        ),
        (
            "Own the football logic",
            "Define how passes, turnovers and shots are detected and reviewed.",
            ORANGE,
        ),
        (
            "Explain every statistic",
            "Keep the supporting video and confidence instead of returning a black-box number.",
            PURPLE,
        ),
    ]
    for index, (title, detail, accent) in enumerate(items):
        row, col = divmod(index, 2)
        add_card(
            slide,
            title,
            detail,
            0.8 + col * 6.15,
            1.35 + row * 2.45,
            5.65,
            2.0,
            accent=accent,
            title_size=21,
            detail_size=16,
        )
    add_text(
        slide,
        "Our long-term advantage is the amateur-football dataset and dependable "
        "camera-to-statistics workflow—not one secret algorithm.",
        1.3,
        6.15,
        10.7,
        0.55,
        size=18,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def progress_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "What we have already proved",
        "A working proof of concept—not yet a finished commercial product",
    )
    add_card(
        slide,
        "Tracking works",
        "Players, teams and the ball are detected and tracked on real panoramic footage.",
        0.75,
        1.35,
        3.75,
        2.25,
        accent=BLUE,
    )
    add_card(
        slide,
        "Football events work",
        "The system produces reviewable pass, turnover and shot candidates.",
        4.8,
        1.35,
        3.75,
        2.25,
        accent=ORANGE,
    )
    add_card(
        slide,
        "The experience works",
        "A browser demo replays video and updates team statistics at event time.",
        8.85,
        1.35,
        3.75,
        2.25,
        accent=GREEN,
    )
    add_text(
        slide,
        "Measured evidence",
        0.8,
        4.35,
        4.2,
        0.4,
        size=21,
        color=PURPLE,
        bold=True,
    )
    add_text(
        slide,
        "• Three independent video windows tested\n"
        "• 71% pass precision across those windows\n"
        "• First correctly matched shot within 0.12 seconds",
        0.8,
        4.9,
        4.75,
        1.35,
        size=17,
    )
    add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        6.25,
        4.25,
        5.9,
        1.85,
        fill=PANEL,
        line=RED,
    )
    add_text(
        slide,
        "Still to prove",
        6.55,
        4.52,
        2.0,
        0.35,
        size=20,
        color=RED,
        bold=True,
    )
    add_text(
        slide,
        "Two live camera feeds • GPU speed • real home-match lighting • operator "
        "workflow • reliable shot outcomes",
        6.55,
        5.02,
        5.25,
        0.65,
        size=16,
    )


def comparison_slide(prs: Presentation):
    slide = new_slide(prs, "Our position in the market")
    headers = ("", "Typical vendor platform", "Our intended product")
    widths = (3.15, 4.45, 4.45)
    x_values = (0.65, 3.8, 8.25)
    for index, header in enumerate(headers):
        add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x_values[index],
            1.25,
            widths[index],
            0.7,
            fill=PANEL_2,
            line=(MUTED, ORANGE, GREEN)[index],
        )
        add_text(
            slide,
            header,
            x_values[index] + 0.1,
            1.45,
            widths[index] - 0.2,
            0.25,
            size=16,
            color=(MUTED, ORANGE, GREEN)[index],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    rows = [
        ("Camera", "Vendor-specific or bundled", "Standard club-owned IP cameras"),
        ("Processing", "Usually cloud/service based", "Local NVIDIA GPU computer"),
        ("Data control", "Inside vendor platform", "Club-controlled video and evidence"),
        ("Statistics", "Polished but black-box", "Provisional, visible and reviewable"),
        ("Target buyer", "Broad market to professional", "Amateur home-ground pilot first"),
    ]
    for row_index, row in enumerate(rows):
        y = 2.12 + row_index * 0.85
        for col_index, value in enumerate(row):
            fill = PANEL if row_index % 2 == 0 else RGBColor(17, 27, 38)
            add_shape(
                slide,
                MSO_SHAPE.RECTANGLE,
                x_values[col_index],
                y,
                widths[col_index],
                0.72,
                fill=fill,
                line=NAVY,
            )
            add_text(
                slide,
                value,
                x_values[col_index] + 0.14,
                y + 0.16,
                widths[col_index] - 0.28,
                0.38,
                size=14,
                bold=col_index == 0,
                color=WHITE if col_index == 0 else (ORANGE if col_index == 1 else GREEN),
                align=PP_ALIGN.LEFT,
                valign=MSO_ANCHOR.MIDDLE,
            )
    add_text(
        slide,
        "We are not trying to outspend professional systems. We are removing the "
        "parts an amateur club cannot afford.",
        1.1,
        6.55,
        11.1,
        0.42,
        size=18,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def roadmap_slide(prs: Presentation):
    slide = new_slide(
        prs,
        "The practical route to a home-match pilot",
        "Prove each expensive assumption before buying the complete setup",
    )
    steps = [
        ("1", "Camera proof", "Record one original 4K minute and measure ball clarity.", BLUE),
        ("2", "GPU proof", "Process two streams at the required speed.", PURPLE),
        ("3", "Closed match", "Run a full home game without public display.", ORANGE),
        ("4", "Operator trial", "Review disputed events and refine the workflow.", GREEN),
        ("5", "Screen pilot", "Publish a small set of provisional statistics.", BLUE),
    ]
    for index, (number, title, detail, accent) in enumerate(steps):
        y = 1.25 + index * 1.05
        add_shape(
            slide,
            MSO_SHAPE.OVAL,
            0.85,
            y,
            0.55,
            0.55,
            fill=accent,
        )
        add_text(
            slide,
            number,
            0.85,
            y + 0.1,
            0.55,
            0.25,
            size=16,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, title, 1.7, y + 0.03, 2.0, 0.3, size=19, bold=True)
        add_text(slide, detail, 3.85, y + 0.03, 7.85, 0.4, size=16, color=MUTED)
        if index < len(steps) - 1:
            add_line(slide, 1.12, y + 0.55, 1.12, y + 1.03, color=MUTED, width=1.5)
    add_badge(slide, "GO / NO-GO AFTER EACH STEP", 4.65, 6.5, 3.8, ORANGE)


def close_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_badge(slide, "THE PROPOSITION", 5.48, 0.55, 2.35, GREEN)
    add_text(
        slide,
        "Provisional live football statistics\nfor clubs priced out of pro tracking",
        1.05,
        1.38,
        11.25,
        1.55,
        size=34,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Own the cameras. Own the processing. Review the evidence. Engage the crowd.",
        1.55,
        3.35,
        10.25,
        0.6,
        size=22,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_badge(slide, "AFFORDABLE", 3.0, 4.55, 1.75, GREEN)
    add_badge(slide, "LOCAL", 4.95, 4.55, 1.4, BLUE)
    add_badge(slide, "TRANSPARENT", 6.55, 4.55, 1.8, PURPLE)
    add_badge(slide, "AMATEUR-FIRST", 8.55, 4.55, 1.9, ORANGE)
    add_text(
        slide,
        "Next proof: one real camera minute and a two-stream GPU benchmark.",
        2.05,
        5.75,
        9.25,
        0.55,
        size=20,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "FootballVideoAnalysisBuilder",
        4.65,
        6.85,
        4.0,
        0.25,
        size=10,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    cover_slide(prs)
    market_slide(prs)
    opportunity_slide(prs)
    proposition_slide(prs)
    setup_slide(prs)
    process_slide(prs)
    output_slide(prs)
    own_code_slide(prs)
    progress_slide(prs)
    comparison_slide(prs)
    roadmap_slide(prs)
    close_slide(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"Product pitch written to {build().resolve()}")
